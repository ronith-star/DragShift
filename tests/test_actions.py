"""Every action, end to end without the screen: all conversions, the PDF tools, each tool
window's main button (driven the way the app drives it), and the Apple Vision actions."""
import os
import sys
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest  # noqa: E402
from PIL import Image  # noqa: E402
from PySide6.QtCore import QPointF, QRectF  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

app = QApplication.instance() or QApplication([])

from dragshift import convert, imaging, tools  # noqa: E402
from dragshift.util import category  # noqa: E402
from test_v11 import make_pdf, photo  # noqa: E402

mac = pytest.mark.skipif(sys.platform != "darwin", reason="Apple Vision is macOS only")


@pytest.fixture
def files(tmp_path):
    ph = photo(tmp_path, size=(1200, 800))                     # JPEG with camera + GPS EXIF
    png = tmp_path / "Graphic.png"
    Image.open(ph).resize((600, 400)).save(png)
    pdf = make_pdf(tmp_path, ph, pages=3)
    from docx import Document
    d = Document()
    d.add_heading("Letter", 1)
    d.add_paragraph("Hello from the DragShift action test.")
    docx = tmp_path / "Letter.docx"
    d.save(docx)
    txt = tmp_path / "Notes.txt"
    txt.write_text("Shopping list\n- coffee\n- limes\n", encoding="utf-8")
    return {"photo": ph, "png": png, "pdf": pdf, "docx": docx, "txt": txt}


def opens(p: Path, target: str):
    """The output is a real file of the expected kind."""
    assert p.exists() and p.stat().st_size > 0, p
    if target in ("PDF",):
        with convert.fitz.open(p) as d:
            assert d.page_count >= 1
    elif target == "DOCX":
        from docx import Document
        assert Document(str(p)).paragraphs
    elif target == "TXT":
        assert p.read_text(encoding="utf-8").strip()
    else:
        with Image.open(p) as im:
            im.load()
            assert im.size[0] > 0


@pytest.mark.parametrize("kind", ["photo", "png", "pdf", "docx", "txt"])
def test_every_conversion(files, kind):
    src = files[kind]
    cat = category(src)
    for target in convert.targets_for(cat, src.suffix):
        outs = convert.convert(src, cat, target)
        assert outs, f"{kind} → {target}: nothing written"
        for o in outs:
            opens(o, target)


def test_images_to_one_pdf(files):
    out = convert.images_to_single_pdf([files["photo"], files["png"]])[0]
    with convert.fitz.open(out) as d:
        assert d.page_count == 2


def test_pdf_tools(files):
    pdf = files["pdf"]
    assert imaging.compress_pdf(pdf).exists()
    assert len(imaging.split_pdf(pdf)) == 3
    with convert.fitz.open(imaging.rotate_pdf(pdf)) as d:
        assert all(pg.rotation == 90 for pg in d)
    with convert.fitz.open(imaging.protect_pdf(pdf, "pw1234")) as d:
        assert d.needs_pass
    assert imaging.extract_pdf_images(pdf)
    assert not imaging.read_pdf_metadata(imaging.strip_pdf_metadata(pdf)).get("author")
    with convert.fitz.open(imaging.merge_pdfs([pdf, pdf])) as d:
        assert d.page_count == 6


class Saved:
    """Stands in for the app's notify(): records what a tool window wrote."""

    def __init__(self):
        self.outs: list[Path] = []

    def __call__(self, action, srcs, outs, extra=""):
        self.outs += [Path(o) for o in outs]


def run_window(cls, src, drive, *args):
    saved = Saved()
    w = cls(*args, saved) if args else cls(str(src) if not isinstance(src, list) else [str(s) for s in src], saved)
    drive(w)
    assert saved.outs, f"{cls.__name__} wrote nothing"
    return saved.outs


def test_compress_window(files):
    out = run_window(tools.CompressWindow, [files["photo"]], lambda w: w.save())[0]
    assert out.stat().st_size < files["photo"].stat().st_size


def test_resize_window(files):
    out = run_window(tools.ResizeWindow, [files["photo"]], lambda w: w.save())[0]
    assert Image.open(out).size[0] > 0


def test_crop_window_square(files):
    def drive(w):
        w.grp.button([n for n, _ in w.RATIOS].index("1:1")).click()
        w.apply()
    out = run_window(tools.CropWindow, files["photo"], drive)[0]
    a, b = Image.open(out).size
    assert abs(a - b) <= 1


def test_edit_window_auto_and_rotate(files):
    def drive(w):
        w.auto()
        w._rot(90)
        w.save()
    out = run_window(tools.EditWindow, files["photo"], drive)[0]
    assert Image.open(out).size == (800, 1200)


def test_annotate_window(files):
    def drive(w):
        w.view.add(tools.Shape("Step", "#ff3b30", 4, [QPointF(60, 60)], "1"))
        w.save()
    out = run_window(tools.AnnotateWindow, files["photo"], drive)[0]
    assert "Annotated" in out.name


def test_redact_window_blacks_out(files):
    def drive(w):
        w.view.rects.append(QRectF(10, 10, 120, 90))
        w.save()
    out = run_window(tools.RedactWindow, files["photo"], drive)[0]
    k = 1200 / Image.open(files["photo"]).size[0]
    px = Image.open(out).convert("RGB").getpixel((int(40 * k), int(40 * k)))
    assert max(px) < 20 and imaging.count_metadata(out) == (0, False)


def test_add_bg_window(files):
    out = run_window(tools.AddBgWindow, files["png"], lambda w: w.save())[0]
    a, b = Image.open(out).size
    assert a > 600 and b > 400


def test_metadata_window(files):
    loc = run_window(tools.MetadataWindow, files["photo"], lambda w: w.strip_location())[0]
    assert imaging.count_metadata(loc)[1] is False and imaging.metadata_summary(loc).get("Camera")
    clean = run_window(tools.MetadataWindow, files["photo"], lambda w: w.strip())[0]
    assert imaging.count_metadata(clean) == (0, False)


def test_protect_window(files):
    def drive(w):
        w.pw.setText("secret12")
        w.pw2.setText("secret12")
        w.save()
    out = run_window(tools.ProtectWindow, files["pdf"], drive)[0]
    with convert.fitz.open(out) as d:
        assert d.needs_pass and d.authenticate("secret12")


@mac
def test_copy_text_through_the_app(files):
    from dragshift.app import DragShiftApp
    r = DragShiftApp.text_result([str(files["pdf"])])
    assert "Page 1" in r.copy and r.title == "Text copied"


@mac
def test_remove_bg_on_a_real_photo(tmp_path):
    from dragshift import platform_mac
    if not platform_mac.can_lift_subject():
        pytest.skip("needs macOS 14")
    sample = Path("/Library/User Pictures/Animals/Penguin.heic")
    if not sample.exists():
        pytest.skip("no sample photo")
    src = tmp_path / "Penguin.heic"
    src.write_bytes(sample.read_bytes())
    out = imaging.remove_background(src)
    with Image.open(out) as im:
        a = im.getchannel("A")
        assert a.getpixel((5, 5)) < 30            # the sky is gone
        assert a.getpixel((330, 450)) > 220       # the penguin's white chest stays
